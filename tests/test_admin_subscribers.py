"""Tests for admin subscriber management routes (list, delete, export)."""
from models.subscriber import Subscriber


def _make_subscriber(db, **overrides):
    defaults = dict(name="Jane Doe", email="jane@example.com")
    defaults.update(overrides)
    subscriber = Subscriber(**defaults)
    db.session.add(subscriber)
    db.session.commit()
    return subscriber


def _xlsx_strings(data: bytes) -> list:
    """Flatten all non-empty cell values from the export's first worksheet."""
    from openpyxl import load_workbook
    import io as _io

    wb = load_workbook(_io.BytesIO(data))
    ws = wb.active
    return [str(cell.value) for row in ws.iter_rows() for cell in row if cell.value is not None]


class TestSubscribersList:
    def test_requires_login(self, client):
        response = client.get("/admin/subscribers")
        assert response.status_code in (302, 401, 403)

    def test_rejects_non_admin(self, app, authenticated_client):
        response = authenticated_client.get("/admin/subscribers")
        assert response.status_code in (302, 403)

    def test_admin_can_access(self, app, admin_client):
        response = admin_client.get("/admin/subscribers")
        assert response.status_code == 200

    def test_lists_existing_subscribers(self, app, admin_client):
        from app import db

        _make_subscriber(db, name="Juan dela Cruz", email="juan@example.com")
        response = admin_client.get("/admin/subscribers")
        assert b"Juan dela Cruz" in response.data
        assert b"juan@example.com" in response.data

    def test_search_by_name(self, app, admin_client):
        from app import db

        _make_subscriber(db, name="Juan dela Cruz", email="juan@example.com")
        _make_subscriber(db, name="Maria Santos", email="maria@example.com")
        response = admin_client.get("/admin/subscribers?search=Juan")
        assert b"Juan dela Cruz" in response.data
        assert b"Maria Santos" not in response.data

    def test_search_by_email(self, app, admin_client):
        from app import db

        _make_subscriber(db, name="Findable", email="findme@example.com")
        _make_subscriber(db, name="Other", email="other@example.com")
        response = admin_client.get("/admin/subscribers?search=findme")
        assert b"findme@example.com" in response.data
        assert b"other@example.com" not in response.data


class TestDeleteSubscriber:
    def test_requires_login(self, client):
        response = client.post("/admin/subscribers/delete/1")
        assert response.status_code in (302, 401, 403)

    def test_rejects_non_admin(self, app, authenticated_client):
        from app import db

        subscriber = _make_subscriber(db)
        authenticated_client.post(f"/admin/subscribers/delete/{subscriber.id}")
        assert db.session.get(Subscriber, subscriber.id) is not None

    def test_admin_can_delete(self, app, admin_client):
        from app import db

        subscriber = _make_subscriber(db)
        subscriber_id = subscriber.id
        response = admin_client.post(f"/admin/subscribers/delete/{subscriber_id}")
        assert response.status_code == 302
        assert db.session.get(Subscriber, subscriber_id) is None

    def test_nonexistent_subscriber_returns_404(self, app, admin_client):
        response = admin_client.post("/admin/subscribers/delete/99999")
        assert response.status_code == 404


class TestExportSubscribers:
    def test_requires_login(self, client):
        response = client.get("/admin/subscribers/export")
        assert response.status_code in (302, 401, 403)

    def test_returns_xlsx(self, app, admin_client):
        from app import db

        _make_subscriber(db)
        response = admin_client.get("/admin/subscribers/export")
        assert response.status_code == 200
        assert "spreadsheetml.sheet" in response.content_type

    def test_header_row(self, app, admin_client):
        response = admin_client.get("/admin/subscribers/export")
        values = _xlsx_strings(response.data)
        assert "Name" in values
        assert "Email" in values
        assert "Subscribed At" in values

    def test_contains_subscriber_data(self, app, admin_client):
        from app import db

        _make_subscriber(db, name="Export Me", email="exportme@example.com")
        response = admin_client.get("/admin/subscribers/export")
        values = _xlsx_strings(response.data)
        assert "Export Me" in values
        assert "exportme@example.com" in values
